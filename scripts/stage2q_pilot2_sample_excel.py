"""재파일럿용 엑셀(pilot2.xlsx) 생성. API 호출 없음, 로컬 데이터만 사용.

정의:
  A_N2_예측, B_N2_예측 = stage2d_gpt_predictions.parquet의 조건 A/B, n==2 예측
  is_mismatch = A_N2_예측 != gold_actual
  is_match    = A_N2_예측 == gold_actual
  음성이_구함  = is_mismatch AND (B_N2_예측 == gold_actual)
  음성이_망침  = is_match    AND (B_N2_예측 != gold_actual)

윈도우: N=2, customer_only (조건 A/B와 매칭). r_gold_valid_flag==True는 파이프라인
전체에서 이미 보장됨. 지난번 발견한 버그(윈도우가 N개로 안 채워진 콜 혼입) 재발
방지를 위해 "윈도우 발화 수 == 2"인 콜만 후보로 인정.

샘플: mismatch 15 + match 15 = 30, 각 층 내부에서 gold_actual 카테고리를
라운드로빈으로 순회하며 뽑아 다양성 확보(강제 비율 없음). 최종 30개 순서 셔플.

출력: outputs/pilot2.xlsx (시트1 "문제지" 보임 / 시트2 "답지" 숨김)
"""

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent.parent
PRED_PATH = BASE_DIR / "outputs" / "stage2d_gpt_predictions.parquet"
GOLD_PATH = BASE_DIR / "outputs" / "gold_actual_batch1_final.parquet"
WINDOWS_NL_PATH = BASE_DIR / "outputs" / "stage2_windows_nl_v2.parquet"
D04_PATH = BASE_DIR / "outputs" / "d04_dialog_index.parquet"
OUT_XLSX = BASE_DIR / "outputs" / "pilot2.xlsx"

SEED = 42
WINDOW_N = 2
WINDOW_VERSION = "customer_only"
N_MISMATCH = 15
N_MATCH = 15


def diverse_sample(candidate_ids: list, gold_map: pd.Series, n_needed: int, seed: int) -> list:
    """gold_actual 카테고리를 라운드로빈으로 순회하며 뽑아 다양성을 확보한다."""
    rng = np.random.default_rng(seed)
    by_cat = {}
    for cid in candidate_ids:
        by_cat.setdefault(gold_map.loc[cid], []).append(cid)
    for cat in by_cat:
        rng.shuffle(by_cat[cat])
    cats = list(by_cat.keys())
    rng.shuffle(cats)

    selected = []
    while len(selected) < n_needed and any(by_cat[c] for c in cats):
        for c in cats:
            if len(selected) >= n_needed:
                break
            if by_cat[c]:
                selected.append(by_cat[c].pop(0))
    return selected


def build_call_table():
    preds = pd.read_parquet(PRED_PATH)
    gold = pd.read_parquet(GOLD_PATH)[["call_id", "label"]].rename(columns={"label": "gold_actual"})

    n2 = preds[(preds["n"] == 2) & (preds["condition"].isin(["A", "B"]))]
    pivot = n2.pivot_table(index="call_id", columns="condition", values="predicted_label", aggfunc="first")
    pivot = pivot.rename(columns={"A": "A_N2_pred", "B": "B_N2_pred"})

    table = gold.merge(pivot, on="call_id", how="inner").dropna(subset=["A_N2_pred", "B_N2_pred"])

    table["is_mismatch"] = table["A_N2_pred"] != table["gold_actual"]
    table["is_match"] = ~table["is_mismatch"]
    table["voice_saved"] = table["is_mismatch"] & (table["B_N2_pred"] == table["gold_actual"])
    table["voice_ruined"] = table["is_match"] & (table["B_N2_pred"] != table["gold_actual"])
    return table


def restrict_to_complete_windows(table: pd.DataFrame, windows: pd.DataFrame) -> pd.DataFrame:
    win_sub = windows[(windows["window_n"] == WINDOW_N) & (windows["window_version"] == WINDOW_VERSION)]
    group_sizes = win_sub.groupby("call_id").size()
    full_ids = set(group_sizes[group_sizes == WINDOW_N].index)
    nl_ok = win_sub.groupby("call_id")["nl_description"].apply(lambda s: s.notna().all())
    nl_ok_ids = set(nl_ok[nl_ok].index)
    valid_ids = full_ids & nl_ok_ids
    return table[table["call_id"].isin(valid_ids)]


def main():
    table = build_call_table()
    windows = pd.read_parquet(WINDOWS_NL_PATH)
    table = restrict_to_complete_windows(table, windows)

    gold_map = table.set_index("call_id")["gold_actual"]

    mismatch_ids = table[table["is_mismatch"]]["call_id"].tolist()
    match_ids = table[table["is_match"]]["call_id"].tolist()

    chosen_mismatch = diverse_sample(mismatch_ids, gold_map, N_MISMATCH, SEED)
    chosen_match = diverse_sample(match_ids, gold_map, N_MATCH, SEED + 1)

    selected_ids = chosen_mismatch + chosen_match
    order_rng = np.random.default_rng(SEED + 2)
    order = np.arange(len(selected_ids))
    order_rng.shuffle(order)
    selected_ids = [selected_ids[i] for i in order]

    # === 검증 출력 ===
    sel_table = table[table["call_id"].isin(selected_ids)].set_index("call_id").loc[selected_ids].reset_index()
    print(f"샘플 총 {len(sel_table)}건 (seed={SEED}, match seed={SEED+1}, 순서셔플 seed={SEED+2})")
    print(f"is_mismatch 합계: {sel_table['is_mismatch'].sum()} (기대값 15)")
    print(f"is_match 합계: {sel_table['is_match'].sum()} (기대값 15)")
    print(f"음성이_구함: {sel_table['voice_saved'].sum()}건")
    print(f"음성이_망침: {sel_table['voice_ruined'].sum()}건")
    print(f"\ngold_actual 분포(참고용, 다양성 확인):")
    print(sel_table["gold_actual"].value_counts().to_string())

    # === 발화/음성서술 텍스트 조립 ===
    win_sub = windows[(windows["window_n"] == WINDOW_N) & (windows["window_version"] == WINDOW_VERSION)]
    subcat = pd.read_parquet(D04_PATH)[["call_id", "subcategory"]].drop_duplicates("call_id").set_index("call_id")

    rows_q, rows_a = [], []
    for idx, cid in enumerate(selected_ids, start=1):
        g = win_sub[win_sub["call_id"] == cid].sort_values("dialog_idx")
        text_lines = [f"발화{i}({r.speaker_group}): {r.text_clean}" for i, r in enumerate(g.itertuples(), 1)]
        nl_lines = [f"발화{i}: {r.nl_description}" for i, r in enumerate(g.itertuples(), 1)]

        rec = sel_table[sel_table["call_id"] == cid].iloc[0]
        rows_q.append({
            "순번": idx, "call_id": cid,
            "대화(텍스트)": "\n".join(text_lines),
            "음성서술": "\n".join(nl_lines),
            "내판단_텍스트만": "", "내판단_음성후": "",
        })
        rows_a.append({
            "순번": idx, "call_id": cid,
            "gold_actual": rec["gold_actual"], "A_N2_예측": rec["A_N2_pred"], "B_N2_예측": rec["B_N2_pred"],
            "is_mismatch": bool(rec["is_mismatch"]), "음성이_구함": bool(rec["voice_saved"]),
            "음성이_망침": bool(rec["voice_ruined"]),
            "subcategory": subcat.loc[cid, "subcategory"] if cid in subcat.index else None,
        })

    build_excel(rows_q, rows_a)
    print(f"\n저장: {OUT_XLSX}")


def build_excel(rows_q, rows_a):
    wb = Workbook()

    # ---- 시트1: 문제지 ----
    ws1 = wb.active
    ws1.title = "문제지"
    headers_q = ["순번", "call_id", "대화(텍스트)", "음성서술", "내판단_텍스트만", "내판단_음성후"]
    ws1.append(headers_q)
    for cell in ws1[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")

    for r in rows_q:
        ws1.append([r[h] for h in headers_q])

    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row):
        for cell in row:
            cell.alignment = wrap

    col_widths = {"A": 6, "B": 16, "C": 55, "D": 55, "E": 22, "F": 22}
    for col, width in col_widths.items():
        ws1.column_dimensions[col].width = width

    for i, r in enumerate(rows_q, start=2):
        n_lines = max(r["대화(텍스트)"].count("\n"), r["음성서술"].count("\n")) + 1
        ws1.row_dimensions[i].height = max(30, n_lines * 15)

    # D열(음성서술) 그룹화(접기/펼치기 가능하게)
    ws1.column_dimensions["D"].outlineLevel = 1
    ws1.sheet_properties.outlinePr.summaryRight = True
    ws1.sheet_view.showOutlineSymbols = True

    ws1.freeze_panes = "A2"

    # ---- 시트2: 답지 (숨김) ----
    ws2 = wb.create_sheet("답지")
    ws2["A1"] = "⚠️ 채점 전 열람 금지"
    ws2["A1"].font = Font(bold=True, color="FF0000")

    headers_a = ["순번", "call_id", "gold_actual", "A_N2_예측", "B_N2_예측",
                 "is_mismatch", "음성이_구함", "음성이_망침", "subcategory"]
    ws2.append([])  # row2 비움 후 row3에 헤더
    ws2.append(headers_a)
    for cell in ws2[3]:
        cell.font = Font(bold=True)

    for r in rows_a:
        ws2.append([r[h] for h in headers_a])

    for col_letter, width in zip("ABCDEFGHI", [6, 16, 12, 12, 12, 12, 12, 12, 16]):
        ws2.column_dimensions[col_letter].width = width

    ws2.sheet_state = "hidden"

    wb.save(OUT_XLSX)


if __name__ == "__main__":
    main()
