"""프롬프트 파일럿 결과 집계 -> outputs/prompt_pilot.xlsx (시트 2개) + 5줄 요약."""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sklearn.metrics import (
    accuracy_score, f1_score, precision_score, recall_score,
)

BASE_DIR = Path(__file__).resolve().parent.parent
PRED_PATH = BASE_DIR / "outputs" / "prompt_pilot_gpt_predictions.parquet"
GOLD_PATH = BASE_DIR / "outputs" / "gold_actual_batch1_final.parquet"
SAMPLE_A_PATH = BASE_DIR / "outputs" / "prompt_pilot_sample_a.csv"
SAMPLE_B_PATH = BASE_DIR / "outputs" / "prompt_pilot_sample_b.csv"
OUT_XLSX = BASE_DIR / "outputs" / "prompt_pilot.xlsx"

CATEGORY_ORDER = ["환불요청", "주문취소", "불만제기", "배송확인", "교환반품", "구매진행", "서비스이용"]
NEUTRAL_LABELS = ["배송확인", "구매진행", "서비스이용", "교환반품"]
REFUND_CANCEL_LABELS = ["환불요청", "주문취소"]


def classify_change(base_pred, other_pred, gold):
    if base_pred == other_pred:
        return "변화없음"
    base_ok = base_pred == gold
    other_ok = other_pred == gold
    if not base_ok and other_ok:
        return "개선"
    if base_ok and not other_ok:
        return "개악"
    return "변화(무관)"


def main():
    preds = pd.read_parquet(PRED_PATH)
    gold = pd.read_parquet(GOLD_PATH)[["call_id", "label"]].rename(columns={"label": "gold_actual"})
    sample_a = pd.read_csv(SAMPLE_A_PATH)
    sample_b = pd.read_csv(SAMPLE_B_PATH)

    print(f"예측 결과: {len(preds)}건, sample x version 분포:")
    print(preds.groupby(["sample", "version"]).size().unstack())

    pivot = preds.pivot_table(index="call_id", columns="version", values="predicted_label", aggfunc="first")

    # === 샘플 A ===
    a = sample_a.merge(pivot[["base", "kr_full"]], on="call_id", how="left")
    a = a.rename(columns={"base": "v_base_예측", "kr_full": "v_kr_full_예측"})
    a["변화"] = a.apply(lambda r: classify_change(r["v_base_예측"], r["v_kr_full_예측"], r["gold_actual"]), axis=1)

    n_null_a = a[["v_base_예측", "v_kr_full_예측"]].isna().sum().sum()
    print(f"\n샘플 A 예측 결측: {n_null_a}건")

    def complaint_metrics(df, pred_col):
        y_true = (df["gold_actual"] == "불만제기").astype(int)
        y_pred = (df[pred_col] == "불만제기").astype(int)
        return {
            "recall": recall_score(y_true, y_pred, zero_division=0),
            "precision": precision_score(y_true, y_pred, zero_division=0),
            "f1": f1_score(y_true, y_pred, zero_division=0),
        }

    a_metrics = {}
    for pred_col, tag in [("v_base_예측", "base"), ("v_kr_full_예측", "kr_full")]:
        m = complaint_metrics(a, pred_col)
        false_alarm = ((a["gold_actual"].isin(REFUND_CANCEL_LABELS)) & (a[pred_col] == "불만제기")).sum()
        neutral_misfire = ((a["gold_actual"].isin(NEUTRAL_LABELS)) & (a[pred_col] == "불만제기")).sum()
        a_metrics[tag] = {**m, "오탐(환불/취소->불만)": false_alarm, "안전성위반(중립->불만)": neutral_misfire}

    n_changed = (a["변화"] != "변화없음").sum()
    change_dir = a["변화"].value_counts().to_dict()

    print("\n=== 샘플 A: base vs kr_full 불만제기 지표 ===")
    print(pd.DataFrame(a_metrics).T.to_string())
    print(f"\n예측 바뀐 콜: {n_changed}건, 방향: {change_dir}")

    # === 샘플 B ===
    b = sample_b.merge(pivot[["base", "kr_slim"]], on="call_id", how="left")
    b = b.rename(columns={"base": "v_base_예측", "kr_slim": "v_kr_slim_예측"})
    b["변화"] = b.apply(lambda r: classify_change(r["v_base_예측"], r["v_kr_slim_예측"], r["gold_actual"]), axis=1)

    n_null_b = b[["v_base_예측", "v_kr_slim_예측"]].isna().sum().sum()
    print(f"\n샘플 B 예측 결측: {n_null_b}건")

    b_overall = {}
    for pred_col, tag in [("v_base_예측", "base"), ("v_kr_slim_예측", "slim")]:
        acc = accuracy_score(b["gold_actual"], b[pred_col])
        macro_f1 = f1_score(b["gold_actual"], b[pred_col], labels=CATEGORY_ORDER, average="macro", zero_division=0)
        b_overall[tag] = {"accuracy": acc, "macro_f1": macro_f1}

    n_changed_b = (b["변화"] != "변화없음").sum()
    change_dir_b = b["변화"].value_counts().to_dict()
    net_effect = change_dir_b.get("개선", 0) - change_dir_b.get("개악", 0)

    print("\n=== 샘플 B: base vs kr_slim 전체 성능 ===")
    print(pd.DataFrame(b_overall).T.to_string())
    print(f"예측 바뀐 콜: {n_changed_b}건, 방향: {change_dir_b}, 순효과(개선-개악)={net_effect:+d}")

    class_f1 = {}
    for pred_col, tag in [("v_base_예측", "base"), ("v_kr_slim_예측", "slim")]:
        f1_per_class = f1_score(b["gold_actual"], b[pred_col], labels=CATEGORY_ORDER, average=None, zero_division=0)
        class_f1[tag] = dict(zip(CATEGORY_ORDER, f1_per_class))
    class_f1_df = pd.DataFrame(class_f1)
    print("\n클래스별 F1 (base vs slim):")
    print(class_f1_df.to_string())

    # === 엑셀 저장 ===
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "A_불만효과"
    cols_a = ["call_id", "층", "gold_actual", "v_base_예측", "v_kr_full_예측", "변화"]
    ws1.append(cols_a)
    for cell in ws1[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
    for _, r in a.iterrows():
        ws1.append([r[c] for c in cols_a])

    ws1.append([])
    ws1.append(["=== 지표 (base vs kr_full) ==="])
    ws1.append(["버전", "recall", "precision", "f1", "오탐(환불/취소->불만)", "안전성위반(중립->불만)"])
    for tag, m in a_metrics.items():
        ws1.append([tag, m["recall"], m["precision"], m["f1"], m["오탐(환불/취소->불만)"], m["안전성위반(중립->불만)"]])
    ws1.append([])
    ws1.append(["예측 바뀐 콜 수", n_changed])
    for k, v in change_dir.items():
        ws1.append([f"  {k}", v])

    ws2 = wb.create_sheet("B_실전성능")
    cols_b = ["call_id", "gold_actual", "is_mismatch", "v_base_예측", "v_kr_slim_예측", "변화"]
    ws2.append(cols_b)
    for cell in ws2[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
    for _, r in b.iterrows():
        ws2.append([r[c] for c in cols_b])

    ws2.append([])
    ws2.append(["=== 전체 성능 (base vs kr_slim) ==="])
    ws2.append(["버전", "accuracy", "macro_f1"])
    for tag, m in b_overall.items():
        ws2.append([tag, m["accuracy"], m["macro_f1"]])
    ws2.append([])
    ws2.append(["예측 바뀐 콜 수", n_changed_b, "순효과(개선-개악)", net_effect])
    for k, v in change_dir_b.items():
        ws2.append([f"  {k}", v])
    ws2.append([])
    ws2.append(["=== 클래스별 F1 ==="])
    ws2.append(["class", "base", "slim"])
    for cls in CATEGORY_ORDER:
        ws2.append([cls, class_f1_df.loc[cls, "base"], class_f1_df.loc[cls, "slim"]])

    for ws in (ws1, ws2):
        for col_cells in ws.columns:
            length = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 10), 40)

    wb.save(OUT_XLSX)
    print(f"\n저장: {OUT_XLSX}")

    # === 5줄 요약 ===
    full_gain_recall = a_metrics["kr_full"]["recall"] - a_metrics["base"]["recall"]
    full_gain_precision = a_metrics["kr_full"]["precision"] - a_metrics["base"]["precision"]
    slim_macro_gain = b_overall["slim"]["macro_f1"] - b_overall["base"]["macro_f1"]
    slim_acc_gain = b_overall["slim"]["accuracy"] - b_overall["base"]["accuracy"]

    print("\n" + "=" * 60)
    print("요약")
    print("=" * 60)
    print(f"1. 샘플A: v_kr_full recall {a_metrics['base']['recall']:.3f}->{a_metrics['kr_full']['recall']:.3f} "
          f"({full_gain_recall:+.3f}), precision {a_metrics['base']['precision']:.3f}->"
          f"{a_metrics['kr_full']['precision']:.3f} ({full_gain_precision:+.3f})")
    print(f"2. 샘플A 안전성: 중립->불만 오분류 base={a_metrics['base']['안전성위반(중립->불만)']}건 -> "
          f"kr_full={a_metrics['kr_full']['안전성위반(중립->불만)']}건, "
          f"환불/취소->불만 오탐 base={a_metrics['base']['오탐(환불/취소->불만)']}건 -> "
          f"kr_full={a_metrics['kr_full']['오탐(환불/취소->불만)']}건")
    print(f"3. 샘플B: accuracy {b_overall['base']['accuracy']:.3f}->{b_overall['slim']['accuracy']:.3f} "
          f"({slim_acc_gain:+.3f}), macro-F1 {b_overall['base']['macro_f1']:.3f}->"
          f"{b_overall['slim']['macro_f1']:.3f} ({slim_macro_gain:+.3f}) "
          f"-> {'채택 권장(순이득 +)' if slim_macro_gain > 0 else '채택 보류(순이득 -)'}")
    print(f"4. 샘플B 예측변화 순효과(개선-개악) = {net_effect:+d}건 (개선{change_dir_b.get('개선',0)}/"
          f"개악{change_dir_b.get('개악',0)}/무관{change_dir_b.get('변화(무관)',0)})")
    print(f"5. 샘플B 자연분포(불만 비율 {b['gold_actual'].value_counts(normalize=True).get('불만제기',0)*100:.1f}%): "
          f"{b['gold_actual'].value_counts().to_dict()}")

    print("\nPROMPT_PILOT_EXCEL_COMPLETE")


if __name__ == "__main__":
    main()
