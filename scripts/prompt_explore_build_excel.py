"""prompt_explore(base/v2/v3/v4) 결과 집계 -> outputs/prompt_explore.xlsx + 5줄 요약."""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sklearn.metrics import accuracy_score, f1_score, precision_score, recall_score

BASE_DIR = Path(__file__).resolve().parent.parent
PRED_PATH = BASE_DIR / "outputs" / "prompt_explore_gpt_predictions.parquet"
BASE_REUSE_PATH = BASE_DIR / "outputs" / "prompt_explore_base_reused.parquet"
GOLD_PATH = BASE_DIR / "outputs" / "gold_actual_batch1_final.parquet"
SAMPLE_B_PATH = BASE_DIR / "outputs" / "prompt_pilot_sample_b.csv"
OUT_XLSX = BASE_DIR / "outputs" / "prompt_explore.xlsx"

CATEGORY_ORDER = ["환불요청", "주문취소", "불만제기", "배송확인", "교환반품", "구매진행", "서비스이용"]
VERSIONS = ["base", "v2", "v3", "v4"]


def classify_change(base_pred, other_pred, gold):
    if base_pred == other_pred:
        return "변화없음"
    base_ok, other_ok = base_pred == gold, other_pred == gold
    if not base_ok and other_ok:
        return "개선"
    if base_ok and not other_ok:
        return "개악"
    return "변화(무관)"


def complaint_metrics(df, pred_col):
    valid = df.dropna(subset=[pred_col])
    y_true = (valid["gold_actual"] == "불만제기").astype(int)
    y_pred = (valid[pred_col] == "불만제기").astype(int)
    return {
        "recall": recall_score(y_true, y_pred, zero_division=0),
        "precision": precision_score(y_true, y_pred, zero_division=0),
        "f1": f1_score(y_true, y_pred, zero_division=0),
    }


def main():
    preds = pd.read_parquet(PRED_PATH)
    base_reused = pd.read_parquet(BASE_REUSE_PATH)
    gold = pd.read_parquet(GOLD_PATH)[["call_id", "label"]].rename(columns={"label": "gold_actual"})
    sample_b = pd.read_csv(SAMPLE_B_PATH)[["call_id"]]

    print(f"신규 예측: {len(preds)}건, version별:")
    print(preds.groupby("version").size())

    pivot = preds.pivot_table(index="call_id", columns="version", values="predicted_label", aggfunc="first")
    parse_err_pivot = preds.pivot_table(index="call_id", columns="version", values="parse_error", aggfunc="first")

    v4_complaint = preds[preds["version"] == "v4"].set_index("call_id")["complaint_tag"]

    table = sample_b.merge(gold, on="call_id", how="left")
    table = table.merge(base_reused, on="call_id", how="left")
    for v in ["v2", "v3", "v4"]:
        table[v] = table["call_id"].map(pivot[v]) if v in pivot.columns else None
    table["v4_불만판단"] = table["call_id"].map(v4_complaint)

    n_base_missing = table["base"].isna().sum()
    print(f"\nv_base 결측(재사용 실패): {n_base_missing}건")

    # === 파싱 실패율 ===
    print("\n=== 버전별 파싱 실패(형식오류 등) 건수 ===")
    parse_fail = {}
    for v in ["v2", "v3", "v4"]:
        if v in parse_err_pivot.columns:
            n_fail = parse_err_pivot[v].notna().sum()
        else:
            n_fail = 0
        parse_fail[v] = n_fail
        print(f"  {v}: {n_fail}/{len(table)} ({n_fail/len(table)*100:.1f}%)")

    # === 버전별 지표 ===
    overall_metrics = {}
    for v in VERSIONS:
        valid = table.dropna(subset=[v])
        acc = accuracy_score(valid["gold_actual"], valid[v]) if len(valid) else float("nan")
        macro_f1 = f1_score(valid["gold_actual"], valid[v], labels=CATEGORY_ORDER, average="macro",
                             zero_division=0) if len(valid) else float("nan")
        cm = complaint_metrics(table, v)
        overall_metrics[v] = {"n_valid": len(valid), "accuracy": acc, "macro_f1": macro_f1, **cm}

    print("\n=== 버전별 전체 성능 ===")
    print(pd.DataFrame(overall_metrics).T.to_string())

    # === 클래스별 F1 ===
    class_f1 = {}
    for v in VERSIONS:
        valid = table.dropna(subset=[v])
        f1s = f1_score(valid["gold_actual"], valid[v], labels=CATEGORY_ORDER, average=None, zero_division=0)
        class_f1[v] = dict(zip(CATEGORY_ORDER, f1s))
    class_f1_df = pd.DataFrame(class_f1)
    print("\n=== 클래스별 F1 ===")
    print(class_f1_df.to_string())

    # === v_base 대비 변화 ===
    change_summary = {}
    for v in ["v2", "v3", "v4"]:
        sub = table.dropna(subset=["base", v])
        sub_change = sub.apply(lambda r: classify_change(r["base"], r[v], r["gold_actual"]), axis=1)
        counts = sub_change.value_counts().to_dict()
        net = counts.get("개선", 0) - counts.get("개악", 0)
        change_summary[v] = {"n": len(sub), **counts, "순효과": net}
        table.loc[sub.index, f"변화_{v}"] = sub_change

    print("\n=== v_base 대비 변화 ===")
    for v, s in change_summary.items():
        print(f"  {v}: {s}")

    # === v4 불만판단 vs gold 불만이진 정확도 ===
    v4_valid = table.dropna(subset=["v4_불만판단"])
    v4_valid = v4_valid[v4_valid["v4_불만판단"].isin(["예", "아니오"])]
    if len(v4_valid):
        y_true = (v4_valid["gold_actual"] == "불만제기").astype(int)
        y_pred = (v4_valid["v4_불만판단"] == "예").astype(int)
        v4_bin_acc = accuracy_score(y_true, y_pred)
        v4_bin_f1 = f1_score(y_true, y_pred, zero_division=0)
        v4_bin_recall = recall_score(y_true, y_pred, zero_division=0)
        v4_bin_precision = precision_score(y_true, y_pred, zero_division=0)
    else:
        v4_bin_acc = v4_bin_f1 = v4_bin_recall = v4_bin_precision = float("nan")
    print(f"\n=== v4 <불만> 태그 vs gold 불만이진 (n={len(v4_valid)}) ===")
    print(f"accuracy={v4_bin_acc:.4f}, F1={v4_bin_f1:.4f}, recall={v4_bin_recall:.4f}, precision={v4_bin_precision:.4f}")

    # === 엑셀 ===
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "콜별결과"
    cols = ["call_id", "gold_actual", "base", "v2", "v3", "v4", "v4_불만판단"]
    ws1.append(cols)
    for cell in ws1[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
    for _, r in table.iterrows():
        ws1.append([r.get(c) for c in cols])

    ws2 = wb.create_sheet("버전별지표")
    ws2.append(["version", "n_valid", "accuracy", "macro_f1", "complaint_recall", "complaint_precision", "complaint_f1"])
    for v in VERSIONS:
        m = overall_metrics[v]
        ws2.append([v, m["n_valid"], m["accuracy"], m["macro_f1"], m["recall"], m["precision"], m["f1"]])
    ws2.append([])
    ws2.append(["클래스별 F1"])
    ws2.append(["class"] + VERSIONS)
    for cls in CATEGORY_ORDER:
        ws2.append([cls] + [class_f1_df.loc[cls, v] for v in VERSIONS])
    ws2.append([])
    ws2.append(["v_base 대비 변화"])
    ws2.append(["version", "n", "변화없음", "개선", "개악", "변화(무관)", "순효과"])
    for v, s in change_summary.items():
        ws2.append([v, s["n"], s.get("변화없음", 0), s.get("개선", 0), s.get("개악", 0),
                    s.get("변화(무관)", 0), s["순효과"]])
    ws2.append([])
    ws2.append(["버전별 파싱 실패"])
    ws2.append(["version", "실패건수", "비율(%)"])
    for v, n_fail in parse_fail.items():
        ws2.append([v, n_fail, round(n_fail / len(table) * 100, 2)])
    ws2.append([])
    ws2.append(["v4 <불만> 태그 vs gold 불만이진"])
    ws2.append(["n", "accuracy", "f1", "recall", "precision"])
    ws2.append([len(v4_valid), v4_bin_acc, v4_bin_f1, v4_bin_recall, v4_bin_precision])

    for ws in (ws1, ws2):
        for col_cells in ws.columns:
            length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 10), 30)

    wb.save(OUT_XLSX)
    print(f"\n저장: {OUT_XLSX}")

    # === 5줄 요약 ===
    print("\n" + "=" * 60)
    print("요약")
    print("=" * 60)
    best_net = max(change_summary.items(), key=lambda kv: kv[1]["순효과"])
    print(f"1. v_base 대비 순효과(개선-개악) 최댓값: {best_net[0]} ({best_net[1]['순효과']:+d}) "
          f"| 전체: " + ", ".join(f"{v}={s['순효과']:+d}" for v, s in change_summary.items()))
    print(f"2. 파싱 안정성: " + ", ".join(f"{v}={parse_fail[v]}/{len(table)}건 실패" for v in ["v2", "v3", "v4"]))
    print(f"3. 전체 macro-F1: " + ", ".join(f"{v}={overall_metrics[v]['macro_f1']:.3f}" for v in VERSIONS))
    print(f"4. 불만제기 recall/precision: " + ", ".join(
        f"{v}=({overall_metrics[v]['recall']:.3f}/{overall_metrics[v]['precision']:.3f})" for v in VERSIONS))
    print(f"5. v4 <불만>태그 단독 이진정확도={v4_bin_acc:.3f} (참고: 최종 <카테고리> 예측과 별개 신호)")

    print("\nPROMPT_EXPLORE_EXCEL_COMPLETE")


if __name__ == "__main__":
    main()
